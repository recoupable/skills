#!/usr/bin/env python3
"""Lossless, local catalog intake. No network calls or hosted writes.

Run with --source, --mapping, --out; --sheet is mandatory for multi-sheet XLSX.
A map names source columns explicitly; the helper never guesses rights or merges
unidentified rows by title. See templates/import-map.json and references/intake.md.
"""
import argparse
import csv
import datetime
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

KINDS = {"composition", "recording", "release", "party"}
INTEREST_KINDS = {"ownership", "writer_interest", "administration", "collection", "statement_percentage", "producer_participation"}
IDENTIFIERS = {"isrc", "iswc", "upc", "ean", "ipi", "isni", "provider_id"}


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def text(value):
    if value is None:
        return ""
    return value.isoformat() if isinstance(value, (datetime.date, datetime.datetime)) else str(value)


def read_rows(path, sheet=None):
    if path.suffix.lower() == ".csv":
        if sheet:
            raise ValueError("--sheet applies only to XLSX")
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = [(i, row, []) for i, row in enumerate(csv.reader(stream), 1)]
        return "csv", rows
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Supported inputs are UTF-8 CSV and XLSX; convert other formats explicitly")
    try:
        import openpyxl
    except ImportError:
        raise ValueError("XLSX intake needs openpyxl — pip3 install openpyxl")
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet is None and len(book.sheetnames) != 1:
            raise ValueError("Choose --sheet explicitly; available: " + ", ".join(book.sheetnames))
        if sheet is not None and sheet not in book.sheetnames:
            raise ValueError("Unknown sheet: " + sheet)
        tab = book[sheet or book.sheetnames[0]]
        rows = []
        for i, cells in enumerate(tab.iter_rows(), 1):
            # Formulas remain raw strings. Numeric identifier cells cannot prove original zeros.
            metadata = [{"column": j, "type": c.data_type, "format": c.number_format}
                        for j, c in enumerate(cells) if c.value is not None]
            rows.append((i, [text(c.value) for c in cells], metadata))
        return tab.title, rows
    finally:
        book.close()


def intake(source, mapping, sheet=None, previous=None):
    source = Path(source)
    source_version = hashlib.sha256(source.read_bytes()).hexdigest()
    sheet_name, rows = read_rows(source, sheet)
    header_row = mapping.get("header_row", 1)
    if not isinstance(header_row, int) or header_row < 1 or header_row > len(rows):
        raise ValueError("header_row must identify an existing row")
    headers = rows[header_row - 1][1]
    if not headers or len(headers) != len(set(headers)) or any(not h for h in headers):
        raise ValueError("Headers must be unique and nonempty; preserve original file and supply an explicit corrected mapping source")
    specs = mapping.get("entities", {})
    if not specs or not set(specs).issubset(KINDS):
        raise ValueError("entities must contain one or more of composition, recording, release, party")
    interest_columns = mapping.get("interest", {})
    membership_columns = mapping.get("membership", {})
    links = mapping.get("recording_work_link", {})
    mapped = [c for spec in specs.values() for c in spec.values()] + list(interest_columns.values()) + list(membership_columns.values()) + list(links.values())
    if any(not isinstance(c, str) or c not in headers for c in mapped):
        raise ValueError("Every mapped value must be an exact source column name")
    transform = digest(mapping)
    import_id = "import_" + digest([source_version, sheet_name, transform])[:24]
    same = bool(previous and previous.get("import_id") == import_id)
    entity_map, row_results, memberships, interests, work_links = {}, [], [], [], []
    entity_rows, identity_conflicts = defaultdict(list), set()
    provider_observations = defaultdict(list)
    for number, values, cells in rows[header_row:]:
        row_id = "row_" + digest([import_id, number])[:24]
        raw = dict(zip(headers, values))
        result = {"row_id": row_id, "locator": {"sheet": sheet_name, "row": number}, "raw_values": values,
                  "raw": raw, "cells": cells, "status": "proposed", "reasons": [], "entity_ids": {}}
        row_results.append(result)
        if not any(values):
            result.update(status="rejected", reasons=["blank source row"])
            continue
        if len(values) != len(headers):
            result["reasons"].append("column count differs from header; all raw values retained")
        for kind, spec in specs.items():
            fields = {field: raw.get(column, "") for field, column in spec.items()}
            if not any(fields.values()):
                continue
            identifiers = {field: value for field, value in fields.items() if field in IDENTIFIERS and value}
            # Provider identity requires both explicit scopes; never infer an account.
            provider_key = None
            unscoped_provider = False
            if identifiers.get("provider_id"):
                missing = [field for field in ("provider_namespace", "account_scope", "provider_id")
                           if not fields.get(field, "").strip()]
                if missing:
                    unscoped_provider = True
                    result["reasons"].append(kind + " provider identity lacks nonblank " + ", ".join(missing))
                else:
                    provider_key = (kind, fields["provider_namespace"], fields["account_scope"], fields["provider_id"])
            key = [kind, sorted(identifiers.items()), fields.get("provider_namespace", "")] if identifiers else [kind, row_id]
            if unscoped_provider:
                key = [kind, row_id]
            elif provider_key:
                key.append(fields["account_scope"])
            entity_id = kind + "_" + digest(key)[:24]
            result["entity_ids"][kind] = entity_id
            if not identifiers:
                result["reasons"].append(kind + " identity remains a row-scoped candidate")
            for cell in cells:
                if cell["column"] < len(headers) and headers[cell["column"]] in [spec[k] for k in IDENTIFIERS if k in spec] and cell["type"] == "n":
                    result["reasons"].append(kind + " has numeric identifier cell; original leading zeros require source review")
            if any(value.startswith("=") for value in fields.values()):
                result["reasons"].append(kind + " contains unevaluated formula")
            entity_rows[entity_id].append(result)
            existing = entity_map.get(entity_id)
            if existing and not provider_key and existing["fields"] != fields:
                identity_conflicts.add(entity_id)
            if not existing:
                entity_map[entity_id] = {"entity_id": entity_id, "kind": kind, "fields": fields, "identity_status": "identifier_candidate" if identifiers and not unscoped_provider else "unresolved", "observations": []}
            entity_map[entity_id]["observations"].append({"row_id": row_id, "source_version": source_version, "fields": fields})
            if provider_key:
                provider_observations[provider_key].append((entity_id, result, fields))
        ids = result["entity_ids"]
        if "release" in ids and "recording" in ids:
            memberships.append({"membership_id": "membership_" + digest([ids["release"], ids["recording"], {k: raw.get(v, "") for k, v in membership_columns.items()}])[:24], "release_id": ids["release"], "recording_id": ids["recording"], "fields": {k: raw.get(v, "") for k, v in membership_columns.items()}, "row_id": row_id})
        if "recording" in ids and "composition" in ids:
            work_links.append({"recording_id": ids["recording"], "composition_id": ids["composition"], "status": "proposed", "fields": {k: raw.get(v, "") for k, v in links.items()}, "row_id": row_id})
        if interest_columns:
            interest = {k: raw.get(v, "") for k, v in interest_columns.items()}
            if any(interest.values()):
                # An edge per observation preserves parented chains. Never sum provider rows.
                interests.append({"interest_id": "interest_" + digest([row_id, interest])[:24], "entity_ids": ids, "fields": interest, "row_id": row_id, "status": "unreviewed"})
                if interest.get("kind") not in INTEREST_KINDS:
                    result["reasons"].append("interest kind missing or unsupported")
                if interest.get("value") and (not interest.get("basis") or not interest.get("scale")):
                    result["reasons"].append("interest percentage lacks basis or scale")
                if interest.get("scale"):
                    try:
                        scale = Decimal(interest["scale"])
                        if not scale.is_finite() or scale <= 0:
                            raise ValueError("nonfinite or nonpositive scale")
                        if interest.get("value"):
                            value = Decimal(interest["value"])
                            if not value.is_finite():
                                raise ValueError("nonfinite value")
                            if value < 0 or value > scale:
                                result["reasons"].append("interest value outside supplied scale; provider semantics require review")
                    except (InvalidOperation, ValueError):
                        result["reasons"].append("interest value/scale requires finite numeric values and a positive scale; source review required")
        if not ids:
            result["reasons"].append("no mapped entity values")
        if result["reasons"]:
            result["status"] = "held"
    for observations in provider_observations.values():
        values = defaultdict(set)
        for _, _, fields in observations:
            for field, value in fields.items():
                if value.strip():
                    values[field].add(value)
        conflicts = sorted(field for field, observed in values.items() if len(observed) > 1)
        if conflicts:
            reason = "conflicting " + ", ".join(conflicts) + " for scoped provider identity"
            for entity_id, row, _ in observations:
                entity_map[entity_id]["identity_status"] = "conflicting_observations"
                row["status"] = "held"
                row["reasons"].append(reason)
    for entity_id in identity_conflicts:
        entity_map[entity_id]["identity_status"] = "conflicting_observations"
        for row in entity_rows[entity_id]:
            row["status"] = "held"
            row["reasons"].append("conflicting metadata for " + entity_id)
    for row in row_results:
        if same and row["status"] == "proposed":
            row["status"] = "unchanged"
    counts = Counter(r["status"] for r in row_results)
    return {"schema_version": "catalog-intake/1", "import_id": import_id, "source": {"filename": source.name, "sha256": source_version, "sheet": sheet_name, "header_row": header_row, "headers": headers, "preamble_rows": rows[:header_row-1]},
            "mapping": mapping, "mapping_sha256": transform, "supersedes_import_id": previous.get("import_id") if previous and not same else None,
            "idempotency": "unchanged" if same else "new_version", "hosted_save": {"status": "not_attempted", "receipt": None},
            "controls": {"input_rows": len(row_results), "accounted_rows": sum(counts.values()), "dispositions": dict(counts), "entity_candidates": dict(Counter(e["kind"] for e in entity_map.values()))},
            "entities": list(entity_map.values()), "release_memberships": memberships, "recording_work_links": work_links, "interest_observations": interests, "rows": row_results}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True)
    parser.add_argument("--mapping", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--previous")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    try:
        output = Path(args.out)
        for filename in (args.source, args.mapping, args.previous):
            if filename:
                original = Path(filename)
                if output.resolve() == original.resolve() or (output.exists() and original.exists() and output.samefile(original)):
                    raise ValueError("Output cannot overwrite the source, mapping, or previous import")
        result = intake(args.source, json.loads(Path(args.mapping).read_text()), args.sheet,
                        json.loads(Path(args.previous).read_text()) if args.previous else None)
    except (ValueError, OSError, csv.Error) as error:
        parser.exit(2, str(error) + "\n")
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result["controls"]))


if __name__ == "__main__":
    main()
